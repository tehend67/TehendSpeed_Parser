import csv
import json
import os
import tempfile

def _rows_to_table(rows: list[dict]) -> list[dict]:
    out = []
    for r in rows:
        out.append({
            "ID": r["id"],
            "Name": r.get("name") or "",
            "Username": f"@{r['username']}" if r.get("username") else "",
            "Premium": "YES" if r.get("is_premium") else "NO",
            "Phone": r.get("phone") or "",
            "Source": r.get("source_chat") or "",
            "Updated": r.get("updated_at") or "",
        })
    return out

def export_csv(rows: list[dict]) -> str:
    fd, path = tempfile.mkstemp(suffix=".csv", prefix="dewflow_")
    os.close(fd)
    table = _rows_to_table(rows)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["ID", "Name", "Username", "Premium", "Phone", "Source", "Updated"])
        w.writeheader()
        w.writerows(table)
    return path

def export_json(rows: list[dict]) -> str:
    fd, path = tempfile.mkstemp(suffix=".json", prefix="dewflow_")
    os.close(fd)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
    return path

def export_txt(rows: list[dict]) -> str:
    fd, path = tempfile.mkstemp(suffix=".txt", prefix="dewflow_")
    os.close(fd)
    with open(path, "w", encoding="utf-8") as f:
        for r in rows:
            uname = f"@{r['username']}" if r.get("username") else "—"
            f.write(f"{r['id']} | {r.get('name','')} | {uname} | {r.get('phone','')} | {r.get('source_chat','')}\n")
    return path

def export_xlsx(rows: list[dict]) -> str:
    import pandas as pd
    fd, path = tempfile.mkstemp(suffix=".xlsx", prefix="dewflow_")
    os.close(fd)
    table = _rows_to_table(rows)
    df = pd.DataFrame(table, columns=["ID", "Name", "Username", "Premium", "Phone", "Source", "Updated"])
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Users")
        ws = writer.sheets["Users"]
        from openpyxl.styles import Font, PatternFill, Alignment
        hdr_fill = PatternFill("solid", fgColor="111827")
        hdr_font = Font(color="FFFFFF", bold=True, size=11)
        for cell in ws[1]:
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 22
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0 for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 45)
    return path
